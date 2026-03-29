#!/usr/bin/env python3
"""
Minnesota Immigration Nonprofits — Digital Presence Audit
Reads ian_mn_organizations.xlsx, visits each org's website,
and collects digital presence, staff, founding year, and service data.

Run:
    cd C:\\Users\\endal\\OneDrive\\Desktop\\PR
    python digital_audit.py

Output: digital_presence_audit.xlsx
"""

# ── Auto-install dependencies ─────────────────────────────────────────────────
import subprocess, sys

for pkg in ["requests", "beautifulsoup4", "openpyxl", "pandas"]:
    try:
        __import__(pkg if pkg != "beautifulsoup4" else "bs4")
    except ImportError:
        print(f"Installing {pkg}...")
        subprocess.check_call([sys.executable, "-m", "pip", "install", pkg, "-q"])

import requests
from bs4 import BeautifulSoup
import pandas as pd
import re
import time
from datetime import date
from urllib.parse import urljoin, urlparse

# ── Config ────────────────────────────────────────────────────────────────────
INPUT_FILE  = "ian_mn_organizations.xlsx"
OUTPUT_FILE = "digital_presence_audit.xlsx"
DELAY       = 2.0   # seconds between requests — be polite

HEADERS = {
    "User-Agent": (
        "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/120.0.0.0 Safari/537.36"
    ),
    "Accept-Language": "en-US,en;q=0.9",
}

# Social media platform patterns
SOCIAL_PATTERNS = {
    "Facebook":  r"facebook\.com/(?!sharer|share|dialog|login|plugins)([A-Za-z0-9._/-]+)",
    "Instagram": r"instagram\.com/([A-Za-z0-9._]+)",
    "Twitter_X": r"(?:twitter|x)\.com/([A-Za-z0-9_]+)",
    "LinkedIn":  r"linkedin\.com/(?:company|in)/([A-Za-z0-9._-]+)",
    "YouTube":   r"youtube\.com/(?:channel|c|user|@)([A-Za-z0-9._-]+)",
}

# Keywords to find staff count, founding year, languages, counties
YEAR_PATTERN     = r"\b(19[5-9]\d|200\d|201\d|202[0-4])\b"
STAFF_PATTERNS   = [
    r"(\d+)\s*(?:full[- ]?time\s*)?staff",
    r"team\s+of\s+(\d+)",
    r"(\d+)\s*employees",
    r"(\d+)\s*attorneys",
]
LANGUAGE_KEYWORDS = [
    "Spanish","Somali","Hmong","Karen","Arabic","French","Vietnamese",
    "Amharic","Tigrinya","Oromo","Swahili","Burmese","Khmer","Lao",
    "Russian","Nepali","Hindi","Portuguese","Dari","Pashto",
]
COUNTY_KEYWORDS = [
    "Hennepin","Ramsey","Anoka","Dakota","Washington","Scott","Carver",
    "Olmsted","St. Louis","Stearns","Sherburne","Wright","Kandiyohi",
    "outstate","greater Minnesota","statewide","all counties","Twin Cities",
    "metro","Rochester","Duluth","St. Cloud","Mankato",
]
SERVICE_KEYWORDS = {
    "Legal-Service":   ["legal representation","immigration attorney","deportation","removal","asylum","immigration court","BIA","USCIS"],
    "Direct-Service":  ["resettlement","case management","job training","English class","ESL","citizenship class","food","housing","workforce"],
    "Advocacy-Oriented":["advocacy","policy","campaign","organize","coalition","lobby","legislative","systemic"],
}

# ── Helper: fetch a page safely ───────────────────────────────────────────────
def fetch(url, timeout=12):
    try:
        if not url.startswith("http"):
            url = "https://" + url
        r = requests.get(url, headers=HEADERS, timeout=timeout, allow_redirects=True)
        r.raise_for_status()
        return r.text, r.url
    except Exception as e:
        return None, str(e)

# ── Helper: extract URL from contact field ────────────────────────────────────
def extract_url(contact_str):
    if not isinstance(contact_str, str):
        return ""
    # Look for http/https URL
    match = re.search(r"https?://[^\s,]+", contact_str)
    if match:
        return match.group(0).rstrip(".,)")
    # Look for bare domain
    match = re.search(r"www\.[^\s,]+", contact_str)
    if match:
        return "https://" + match.group(0).rstrip(".,)")
    return ""

# ── Helper: find social media handles in page HTML ───────────────────────────
def find_social_links(html):
    results = {}
    if not html:
        return results
    for platform, pattern in SOCIAL_PATTERNS.items():
        matches = re.findall(pattern, html, re.IGNORECASE)
        # Clean up and deduplicate
        handles = list(dict.fromkeys(
            m.strip("/").split("?")[0].split("#")[0]
            for m in matches
            if m and len(m) > 1 and "sharer" not in m
        ))
        results[platform + "_Handle"] = handles[0] if handles else ""
        results[platform + "_URL"]    = (
            f"https://facebook.com/{handles[0]}"   if platform == "Facebook"  and handles else
            f"https://instagram.com/{handles[0]}"  if platform == "Instagram" and handles else
            f"https://x.com/{handles[0]}"          if platform == "Twitter_X" and handles else
            f"https://linkedin.com/company/{handles[0]}" if platform == "LinkedIn" and handles else
            f"https://youtube.com/@{handles[0]}"   if platform == "YouTube"   and handles else ""
        )
        results[platform + "_Present"] = "Y" if handles else "N"
    return results

# ── Helper: extract year founded ─────────────────────────────────────────────
def find_founding_year(text):
    # Look for "founded in YEAR", "established YEAR", "since YEAR"
    patterns = [
        r"(?:founded|established|incorporated|created|started|since)\s+(?:in\s+)?(\d{4})",
        r"since\s+(\d{4})",
        r"(\d{4})\s+(?:to present|to today)",
    ]
    for p in patterns:
        m = re.search(p, text, re.IGNORECASE)
        if m:
            yr = int(m.group(1))
            if 1950 <= yr <= 2024:
                return str(yr)
    return ""

# ── Helper: extract staff count ───────────────────────────────────────────────
def find_staff_count(text):
    for p in STAFF_PATTERNS:
        m = re.search(p, text, re.IGNORECASE)
        if m:
            return m.group(1)
    return ""

# ── Helper: detect languages ─────────────────────────────────────────────────
def find_languages(text):
    found = [lang for lang in LANGUAGE_KEYWORDS if lang.lower() in text.lower()]
    return "; ".join(found) if found else ""

# ── Helper: detect counties/regions ─────────────────────────────────────────
def find_counties(text):
    found = [c for c in COUNTY_KEYWORDS if c.lower() in text.lower()]
    return "; ".join(found) if found else ""

# ── Helper: classify service area ────────────────────────────────────────────
def classify_service_area(text):
    scores = {}
    for stype, keywords in SERVICE_KEYWORDS.items():
        scores[stype] = sum(1 for kw in keywords if kw.lower() in text.lower())
    best = max(scores, key=scores.get)
    return best if scores[best] > 0 else "Unknown"

# ── Helper: also check About/Staff page ──────────────────────────────────────
def try_about_page(base_url, html):
    """Try to fetch /about or /staff page for more detail."""
    if not html:
        return ""
    soup = BeautifulSoup(html, "html.parser")
    for a in soup.find_all("a", href=True):
        href = a["href"].lower()
        text = a.get_text(strip=True).lower()
        if any(kw in href or kw in text for kw in ["about","staff","team","who we are"]):
            about_url = urljoin(base_url, a["href"])
            about_html, _ = fetch(about_url)
            if about_html:
                return about_html
    return ""

# ── Main audit function ───────────────────────────────────────────────────────
def audit_org(row):
    name = row.get("Name", "Unknown")
    contact = str(row.get("Contact", ""))
    website_url = extract_url(contact)

    result = {
        "Name":              name,
        "IAN_Location":      row.get("Location", ""),
        "IAN_Contact":       contact,
        "IAN_Legal_Areas":   row.get("Areas of legal assistance", ""),
        "IAN_Service_Types": row.get("Types of legal assistance", ""),
        "Website_URL":       website_url,
        "Website_Live":      "",
        "Facebook_Present":  "N",
        "Facebook_Handle":   "",
        "Facebook_URL":      "",
        "Instagram_Present": "N",
        "Instagram_Handle":  "",
        "Instagram_URL":     "",
        "Twitter_X_Present": "N",
        "Twitter_X_Handle":  "",
        "Twitter_X_URL":     "",
        "LinkedIn_Present":  "N",
        "LinkedIn_Handle":   "",
        "LinkedIn_URL":      "",
        "YouTube_Present":   "N",
        "YouTube_Handle":    "",
        "YouTube_URL":       "",
        "Total_Platforms":   0,
        "Year_Founded":      "",
        "Staff_Count":       "",
        "Languages_Served":  "",
        "Counties_Served":   "",
        "Service_Area_Type": "",
        "Notes":             "",
    }

    if not website_url:
        result["Website_Live"] = "No URL found"
        result["Notes"] = "No website URL in IAN directory"
        return result

    print(f"    Fetching: {website_url}")
    html, final_url = fetch(website_url)

    if not html:
        result["Website_Live"] = "N"
        result["Notes"] = f"Fetch failed: {final_url}"
        return result

    result["Website_Live"] = "Y"
    result["Website_URL"]  = final_url  # update to final URL after redirects

    # Full text for analysis
    soup = BeautifulSoup(html, "html.parser")
    page_text = soup.get_text(" ", strip=True)

    # Social media from homepage
    social = find_social_links(html)
    for k, v in social.items():
        if k in result:
            result[k] = v

    # Try About/Staff page for additional detail
    time.sleep(0.5)
    about_html = try_about_page(final_url, html)
    combined_text = page_text
    if about_html:
        about_soup = BeautifulSoup(about_html, "html.parser")
        combined_text += " " + about_soup.get_text(" ", strip=True)
        # Also check About page for social links
        about_social = find_social_links(about_html)
        for k, v in about_social.items():
            if k in result and not result[k]:  # only fill if not already found
                result[k] = v

    # Extract info from combined text
    result["Year_Founded"]     = find_founding_year(combined_text)
    result["Staff_Count"]      = find_staff_count(combined_text)
    result["Languages_Served"] = find_languages(combined_text)
    result["Counties_Served"]  = find_counties(combined_text)
    result["Service_Area_Type"]= classify_service_area(combined_text)

    # Count total platforms present
    platforms = ["Facebook","Instagram","Twitter_X","LinkedIn","YouTube"]
    result["Total_Platforms"] = sum(
        1 for p in platforms if result.get(p + "_Present") == "Y"
    )

    return result

# ── Save to Excel with formatting ─────────────────────────────────────────────
def save_excel(records):
    df = pd.DataFrame(records)

    # Add sampling frame columns
    df.insert(0, "Date Audited", str(date.today()))
    df["Twin Cities Based (Y/N)"]        = ""
    df["In Study Sample (Y/N)"]          = ""
    df["Exclusion Reason"]               = ""
    df["Facebook_Followers (manual)"]    = ""
    df["Instagram_Followers (manual)"]   = ""
    df["Twitter_Followers (manual)"]     = ""

    with pd.ExcelWriter(OUTPUT_FILE, engine="openpyxl") as writer:
        df.to_excel(writer, index=False, sheet_name="Digital Presence Audit")
        ws = writer.sheets["Digital Presence Audit"]
        ws.freeze_panes = "B2"

        # Color-code platform present columns
        from openpyxl.styles import PatternFill, Font
        green = PatternFill("solid", fgColor="E8F5E9")
        red   = PatternFill("solid", fgColor="FFEBEE")

        platform_cols = {}
        for col in ws.iter_cols(1, ws.max_column, 1, 1):
            if "_Present" in str(col[0].value):
                platform_cols[col[0].column] = col[0].value

        for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
            for cell in row:
                if cell.column in platform_cols:
                    if cell.value == "Y":
                        cell.fill = green
                    elif cell.value == "N":
                        cell.fill = red

        # Auto-width columns
        for col in ws.columns:
            max_len = max((len(str(c.value or "")) for c in col), default=10)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 3, 55)

    print(f"\n  Saved: {OUTPUT_FILE}")

# ── Main ──────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    print("\n" + "="*60)
    print("  Minnesota Immigration Nonprofits — Digital Presence Audit")
    print("="*60)

    # Load IAN data
    try:
        df_ian = pd.read_excel(INPUT_FILE)
        print(f"\n  Loaded {len(df_ian)} organizations from {INPUT_FILE}")
    except FileNotFoundError:
        print(f"\n  ERROR: {INPUT_FILE} not found.")
        print("  Make sure you ran scrape_ian_mn.py first.")
        sys.exit(1)

    records = []
    total = len(df_ian)

    for i, (_, row) in enumerate(df_ian.iterrows(), 1):
        name = row.get("Name", f"Row {i}")
        print(f"\n  [{i}/{total}] {name}")
        result = audit_org(row)
        records.append(result)

        # Progress summary
        platforms_found = result["Total_Platforms"]
        print(f"    Website: {result['Website_Live']} | "
              f"FB: {result['Facebook_Present']} | "
              f"IG: {result['Instagram_Present']} | "
              f"X: {result['Twitter_X_Present']} | "
              f"Platforms: {platforms_found}")

        time.sleep(DELAY)

    print(f"\n  Audit complete. {total} organizations processed.")

    # Summary stats
    df_out = pd.DataFrame(records)
    print("\n  ── Summary ──────────────────────────────────")
    print(f"  Organizations with website:   {(df_out['Website_Live']=='Y').sum()}/{total}")
    print(f"  Organizations with Facebook:  {(df_out['Facebook_Present']=='Y').sum()}/{total}")
    print(f"  Organizations with Instagram: {(df_out['Instagram_Present']=='Y').sum()}/{total}")
    print(f"  Organizations with Twitter/X: {(df_out['Twitter_X_Present']=='Y').sum()}/{total}")
    print(f"  Organizations with LinkedIn:  {(df_out['LinkedIn_Present']=='Y').sum()}/{total}")
    print(f"  Avg platforms per org:        {df_out['Total_Platforms'].mean():.1f}")
    print("  ─────────────────────────────────────────────")

    print("\n  Saving Excel output...")
    save_excel(records)

    print("\n  NOTE: Follower counts require manual verification.")
    print("  Open digital_presence_audit.xlsx and fill in:")
    print("    - Facebook_Followers (manual)")
    print("    - Instagram_Followers (manual)")
    print("    - Twitter_Followers (manual)")
    print("    - Twin Cities Based (Y/N)")
    print("    - In Study Sample (Y/N)")
    print("    - Exclusion Reason")
    print("\n" + "="*60 + "\n")
